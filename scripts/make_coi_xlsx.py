"""Generate public/downloads/coi-tracking-spreadsheet.xlsx. Run: python3 scripts/make_coi_xlsx.py"""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

H = ["Sub name","Trade","Contact","Agent email","GL carrier","GL limit","GL expiry","WC expiry","Auto expiry",
     "Additional insured (Y/N)","Endorsement on file (Y/N)","W-9 on file (Y/N)","License #","License expiry","Status"]
N = 500
wb = Workbook(); ws = wb.active; ws.title = "COI Tracker"
ws.append(H)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1D1A17")
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 32
t = date.today()
ex = [
  ["EXAMPLE – Acme Electric","Electrical","Dana Lee, 555-0100","agent@example.com","Example Mutual",1000000,t+timedelta(200),t+timedelta(150),t+timedelta(300),"Y","Y","Y","EL-12345",t+timedelta(400)],
  ["EXAMPLE – Bright Plumbing","Plumbing","Sam Ortiz, 555-0101","agent2@example.com","Sample Casualty",2000000,t+timedelta(20),t+timedelta(90),None,"Y","N","Y","PL-67890",t+timedelta(365)],
  ["EXAMPLE – Cedar Framing","Framing","Kim Park, 555-0102","agent3@example.com","Demo Insurance Co.",1000000,t-timedelta(10),t+timedelta(60),t+timedelta(60),"N","N","N","",None],
]
for r in ex: ws.append(r)
for r in range(2, N + 2):
    ws.cell(r, 15).value = (f'=IF(COUNT(G{r}:I{r},N{r})=0,"",IF(MIN(G{r}:I{r},N{r})<TODAY(),"Expired",'
                            f'IF(MIN(G{r}:I{r},N{r})<=TODAY()+30,"Expiring in 30 days","OK")))')
    for col in (7, 8, 9, 14): ws.cell(r, col).number_format = "yyyy-mm-dd"
    ws.cell(r, 6).number_format = '"$"#,##0'
for r in range(2, 5):
    for c in range(1, 15): ws.cell(r, c).font = Font(italic=True, color="6B635A")
dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True, showErrorMessage=True,
                    error="Enter Y or N", errorTitle="Y/N only")
ws.add_data_validation(dv); dv.add(f"J2:L{N+1}")
dd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True)
ws.add_data_validation(dd); dd.add(f"G2:I{N+1}"); dd.add(f"N2:N{N+1}")
rng = f"O2:O{N+1}"
for txt, bg, fg in (("Expired","FFC7CE","9C0006"),("Expiring in 30 days","FFEB9C","7A4F00"),("OK","C6EFCE","006100")):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$O2="{txt}"'],
        fill=PatternFill("solid", bgColor=bg), font=Font(color=fg, bold=True)))
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:O{N+1}"
for i, w in enumerate([26,14,22,24,20,12,12,12,12,14,14,12,14,14,20], 1):
    ws.column_dimensions[ws.cell(1, i).column_letter].width = w

ins = wb.create_sheet("Instructions", 0)
lines = ["COI Tracking Spreadsheet – Instructions","",
 "1. Go to the 'COI Tracker' sheet. Rows starting with 'EXAMPLE –' are sample data; delete them before use.",
 "2. Add one row per subcontractor. Enter dates as dates (e.g. 2026-12-31).",
 "3. Y/N columns accept only Y or N.",
 "4. Status is calculated automatically from the earliest of GL, WC, Auto and License expiry vs today's date:",
 "     Expired (red) – earliest date is in the past",
 "     Expiring in 30 days (amber) – earliest date is within the next 30 days",
 "     OK (green) – all entered dates are more than 30 days out. Blank if no dates are entered.",
 "5. Status reflects only the dates you enter. It does not confirm coverage; check with the issuing agent or insurer.",
 "","Want reminders and requirement checks done for you? See SubVault: https://jgividen.github.io/subvault/"]
for l in lines: ins.append([l])
ins["A1"].font = Font(bold=True, size=14); ins.column_dimensions["A"].width = 110
wb.active = 1
wb.save("public/downloads/coi-tracking-spreadsheet.xlsx")
